"""
Build Price_Elasticity_Findings.pptx from the executed notebook's artefacts
(output/findings.json + output/fig_*.png).

Run AFTER executing price_elasticity_analysis.ipynb:
    python build_deck.py
"""
import json
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE = __file__.rsplit("/", 1)[0]
OUT = f"{HERE}/output"
findings = json.load(open(f"{OUT}/findings.json"))
M = findings["meta"]
PORT = findings["portfolio"]

# ---------------------------------------------------------------- palette ---
NAVY   = RGBColor(0x16, 0x23, 0x3A)
INK    = RGBColor(0x1A, 0x24, 0x20)
OCHRE  = RGBColor(0xC9, 0x7A, 0x25)
PAPER  = RGBColor(0xEE, 0xEF, 0xE7)
CARD   = RGBColor(0xFC, 0xFC, 0xF9)
GREY   = RGBColor(0x5B, 0x61, 0x58)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
BLUE   = RGBColor(0x2A, 0x78, 0xD6)
RED    = RGBColor(0xE3, 0x49, 0x48)
HAIR   = RGBColor(0xD8, 0xD9, 0xCC)

FONT = "Calibri"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height


def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = PAPER
    return s


def box(s, x, y, w, h):
    return s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame


def set_text(tf, runs, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT,
             space_after=6, line=1.15):
    tf.word_wrap = True
    first = True
    items = runs if isinstance(runs, list) else [runs]
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(space_after)
        p.line_spacing = line
        if isinstance(item, tuple):
            txt, kw = item
        else:
            txt, kw = item, {}
        r = p.add_run()
        r.text = txt
        r.font.name = FONT
        r.font.size = Pt(kw.get("size", size))
        r.font.bold = kw.get("bold", bold)
        r.font.color.rgb = kw.get("color", color)
    return tf


def header(s, kicker, title):
    t = box(s, 0.6, 0.45, 12.1, 0.4)
    set_text(t, (kicker.upper(), {"size": 12, "color": OCHRE, "bold": True}))
    t2 = box(s, 0.6, 0.85, 12.1, 1.0)
    set_text(t2, (title, {"size": 28, "color": NAVY, "bold": True}), line=1.05)
    ln = s.shapes.add_shape(1, Inches(0.62), Inches(1.72), Inches(1.1), Pt(3))
    ln.fill.solid(); ln.fill.fore_color.rgb = OCHRE; ln.line.fill.background()


def bullets(s, items, x=0.7, y=2.1, w=12.0, h=4.8, size=18, gap=10):
    tf = box(s, x, y, w, h)
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.line_spacing = 1.15
        lvl = 0
        txt = it
        if isinstance(it, tuple):
            txt, lvl = it
        r = p.add_run()
        r.text = ("•  " if lvl == 0 else "–  ") + txt
        r.font.name = FONT
        r.font.size = Pt(size if lvl == 0 else size - 3)
        r.font.color.rgb = INK if lvl == 0 else GREY
        p.level = lvl
    return tf


def table(s, rows, x, y, w, h, col_w=None, header_fill=NAVY, fs=12, first_col_left=True):
    nr, nc = len(rows), len(rows[0])
    gf = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(h))
    tb = gf.table
    tb.first_row = True
    if col_w:
        for i, cw in enumerate(col_w):
            tb.columns[i].width = Inches(cw)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = tb.cell(ri, ci)
            c.margin_left = Inches(0.08)
            c.margin_right = Inches(0.08)
            c.margin_top = Inches(0.03)
            c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tfc = c.text_frame
            tfc.word_wrap = True
            p = tfc.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if (ci == 0 and first_col_left) else PP_ALIGN.CENTER
            r = p.add_run()
            r.text = str(val)
            r.font.name = FONT
            r.font.size = Pt(fs)
            if ri == 0:
                r.font.bold = True
                r.font.color.rgb = WHITE
                c.fill.solid(); c.fill.fore_color.rgb = header_fill
            else:
                r.font.color.rgb = INK
                c.fill.solid()
                c.fill.fore_color.rgb = CARD if ri % 2 else RGBColor(0xF2, 0xF3, 0xEC)
    return tb


def money(v):
    return f"${v:,.0f}"


# ============================================================ 1 · TITLE ===
s = slide()
band = s.shapes.add_shape(1, 0, 0, SW, Inches(0.22))
band.fill.solid(); band.fill.fore_color.rgb = OCHRE; band.line.fill.background()
t = box(s, 0.9, 2.35, 11.5, 2.2)
set_text(t, [
    ("SUPPLY BRAIN  ·  CHARTER PHASE 3 DELIVERABLE", {"size": 14, "color": OCHRE, "bold": True}),
    ("Price Elasticity & Buffer Strategy Routing", {"size": 40, "color": NAVY, "bold": True}),
    ("Findings, method, and the price–inventory exchange rate", {"size": 20, "color": GREY}),
], space_after=10)
f = box(s, 0.9, 6.3, 11.5, 0.8)
set_text(f, (f"Rerun on repo commit {M['commit']} (origin/main, pulled {M['pulled']})   |   "
             f"{M['n_orders_2025']:,} 2025 order lines · {M['n_skus']} SKUs   |   prepared by N. Sunnel, 2 Sep 2026",
             {"size": 12, "color": GREY}))

# ==================================================== 2 · EXEC SUMMARY ===
s = slide()
header(s, "Executive summary", "Four findings")
bullets(s, [
    "All five categories are price-elastic (|β₁| from 1.06 to 3.60). Frozen Foods and "
    "Snacks & Confectionery are highly elastic; Beverages is the least.",
    "Elasticity must be read at segment level, not category — Dairy & Bakery splits from elastic "
    "(Premium ≈ −1.7) to statistically flat (Value ≈ −0.3).",
    "The price–inventory exchange rate favours INVENTORY for four of five categories; "
    "Frozen Foods is the only break-even case (1.23×).",
    "Portfolio-wide exchange rate ≈ 0.33× — $1 of discount-driven buffering offsets only "
    "~33¢ of inventory carrying cost, at MEIO’s safety-stock levels.",
], y=1.95, h=3.5, gap=11, size=16)
card = s.shapes.add_shape(1, Inches(0.7), Inches(5.85), Inches(11.9), Inches(1.2))
card.fill.solid(); card.fill.fore_color.rgb = NAVY; card.line.fill.background()
tf = card.text_frame; tf.word_wrap = True
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
set_text(tf, [
    ("Portfolio exchange rate   0.33×", {"size": 22, "color": WHITE, "bold": True}),
    (f"inventory carrying cost {money(PORT['inv_annual_cost_usd'])}/yr   vs   "
     f"equivalent price-buffer cost {money(PORT['price_buffer_cost_usd'])}/yr", {"size": 13, "color": RGBColor(0xCD,0xD3,0xC4)}),
], align=PP_ALIGN.CENTER, space_after=4)

# ==================================================== 3 · CHARTER ASK ===
s = slide()
header(s, "Scope", "What the charter asked for")
q = s.shapes.add_shape(1, Inches(0.7), Inches(2.0), Inches(11.9), Inches(2.1))
q.fill.solid(); q.fill.fore_color.rgb = NAVY; q.line.fill.background()
tf = q.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
tf.margin_left = Inches(0.4); tf.margin_right = Inches(0.4)
set_text(tf, ("“For each segment, estimate price elasticity and compare the cost of buffering with "
              "inventory vs. buffering with price… The system quantifies the exchange rate: "
              "$1 of price flexibility = $Y of safety stock — a comparison no existing supply "
              "chain system makes systematically at SKU level.”",
              {"size": 16, "color": WHITE}), line=1.3)
bullets(s, [
    "How does demand respond to price — by category, and by finer segment?",
    "Is that response in line with published FMCG norms (a check that the model is sound)?",
    "Per segment: what does $1 of price flexibility buy you in avoided safety-stock cost?",
], y=4.5, h=2.5, gap=12, size=17)

# ==================================================== 4 · METHOD ===
s = slide()
header(s, "Method", "The elasticity model")
bullets(s, [
    "Constant-elasticity (log–log) demand curve with two-way fixed effects:",
    ("ln(qty) = β₀ + β₁·ln(base_price) + β₂·ln(1 − discount) + "
     "α(sku) + γ(period) + ε", 1),
    "β₁ = elasticity to shelf/list price   ·   β₂ = response to promo depth "
    "(negative ⇒ discounts lift volume).",
    "SKU and ISO-week fixed effects are demeaned out first, so the coefficient is within-SKU, "
    "within-week price response — not the fact that cheap SKUs sell in bigger unit counts.",
    "Estimator reproduces segmentation_v2.py §2 exactly, then adds HC1 robust standard errors "
    "for confidence intervals.",
    "Exchange rate: price the MEIO-recommended safety stock two ways —",
    ("inventory cost = SS$ × (holding% + capital%)   |   price cost = |ΔQ% ÷ elasticity| "
     "× avg_price × annual_volume", 1),
    ("exchange rate = inventory cost ÷ price cost      (<1 ⇒ inventory cheaper)", 1),
], y=2.0, gap=10, size=16)

# ==================================================== 5 · DATA ===
s = slide()
header(s, "Inputs", "Data & provenance")
table(s, [
    ["Item", "Value"],
    ["Repo commit", f"{M['commit']}  (origin/main, pulled {M['pulled']})"],
    ["Data window", f"{M['data_window'][0]}  →  {M['data_window'][1]}  (2025 actuals only)"],
    ["Order lines", f"{M['n_orders_2025']:,}"],
    ["SKUs / categories / price tiers", f"{M['n_skus']} / 5 / 2"],
    ["Lines carrying a discount", f"{M['promo_share']*100:.1f}%   (max depth 35%)"],
    ["Elasticity inputs", "orders.csv, sku_master.csv"],
    ["Exchange-rate inputs", "cost_parameters.csv, models/MEIO/output/meio_summary_by_category.csv"],
    ["Notebook", "Price Elasticity/price_elasticity_analysis.ipynb"],
], x=0.7, y=2.1, w=11.9, h=4.0, col_w=[3.7, 8.2], fs=13)
b = box(s, 0.7, 6.4, 11.9, 0.7)
set_text(b, ("Only ~5.7% of lines carry price variation, so β₁ leans on modest within-SKU "
             "list-price drift — a real precision constraint (see Caveats).", {"size": 12, "color": GREY}))

# ==================================================== 6 · CATEGORY FIG ===
s = slide()
header(s, "Finding 1", "Category-level price elasticity")
s.shapes.add_picture(f"{OUT}/fig_category_elasticity.png", Inches(0.55), Inches(1.95), width=Inches(8.4))
bullets(s, [
    "Frozen Foods & Snacks ≈ −3.5: shoppers switch sharply on price.",
    "Beverages least elastic (−1.06) — sits in the premium band; portfolio skews to "
    "differentiated SKUs (spirits, specialty tea), not commodity drinks.",
    "Grocery is textbook mainstream FMCG (−1.83).",
    "All coefficients negative, p ≈ 0, tight CIs.",
], x=9.1, y=2.1, w=3.9, gap=10, size=13)

# ============================================ 7 · CATEGORY REG DETAIL ===
s = slide()
header(s, "Finding 1 · detail", "Regression output by category")
rows = [["Category", "β₁ list price", "SE", "95% CI", "β₂ promo", "R²", "n"]]
for r in sorted(findings["category_elasticity"], key=lambda d: d["elasticity_base"]):
    rows.append([
        r["label"],
        f"{r['elasticity_base']:.2f}",
        f"{r['se_base']:.3f}",
        f"[{r['ci_base_lo']:.2f}, {r['ci_base_hi']:.2f}]",
        f"{r['elasticity_promo']:.2f}",
        f"{r['r2']:.3f}",
        f"{r['n_obs']:,}",
    ])
table(s, rows, x=0.45, y=2.2, w=12.45, h=3.2,
      col_w=[2.8, 1.6, 1.15, 2.5, 1.45, 1.35, 1.6], fs=12)
b = box(s, 0.7, 5.7, 11.9, 1.3)
set_text(b, [
    ("R² is low across the board (0.008–0.062): price explains only a sliver of order-to-order "
     "quantity variation on this synthetic data.", {"size": 13, "color": GREY}),
    ("The reliable outputs are the sign, the ordering, and the rough magnitude — not the third decimal.",
     {"size": 13, "color": GREY}),
], space_after=4)

# ==================================================== 8 · TIER FIG ===
s = slide()
header(s, "Finding 2", "Segment split: category × price tier")
s.shapes.add_picture(f"{OUT}/fig_tier_elasticity.png", Inches(0.55), Inches(1.95), width=Inches(8.4))
bullets(s, [
    "Dairy & Bakery is the case segmentation exists for:",
    ("Premium ≈ −1.7 (genuinely elastic)", 1),
    ("Value ≈ −0.29, p = 0.21 — not distinguishable from zero", 1),
    "A blanket category discount would waste margin on Value SKUs that don’t respond.",
    "Frozen & Snacks are uniformly elastic in both tiers; Beverages shows a milder split.",
], x=9.1, y=2.1, w=3.9, gap=9, size=13)

# ==================================================== 9 · BENCHMARK ===
s = slide()
header(s, "Validation", "Benchmark sanity-check")
rows = [["Category", "β₁ (base)", "Nearest published band", "In a published range?"]]
for r in sorted(findings["benchmark_check"], key=lambda d: d["elasticity_base"]):
    rows.append([r["label"], f"{r['elasticity_base']:.2f}",
                 r["nearest benchmark tier"],
                 "Yes" if r["within a published range?"] else "No"])
table(s, rows, x=0.7, y=2.2, w=11.9, h=2.9, col_w=[3.0, 1.6, 4.6, 2.7], fs=13)
bullets(s, [
    "Elasticity is derived from real discount variation, not assumed — but still cross-checked "
    "against published FMCG ranges.",
    "Every category lands inside a recognised band. No wrong signs, no implausible magnitudes "
    "(|β| > 5), no positive coefficients. The model behaves.",
], y=5.4, h=1.9, gap=10, size=15)

# ==================================================== 10 · XCHG CONCEPT ===
s = slide()
header(s, "Finding 3 · concept", "The price–inventory exchange rate")
bullets(s, [
    "Question: for each category, is it cheaper to absorb demand uncertainty with safety stock, "
    "or with price flexibility (a discount that pulls demand in on demand)?",
    "Fixed quantity of protection = the safety stock MEIO already recommends (nothing fabricated).",
    "Inventory cost  =  optimized SS$  ×  (holding cost % + cost of capital %)",
    "Price cost  =  discount depth needed  ×  average price  ×  annual category volume",
    ("where discount depth needed = | (SS in units / annual volume)  ÷  elasticity |", 1),
    "Exchange rate  =  inventory cost  ÷  price cost",
    ("< 1  →  inventory is the cheaper buffer        > 1  →  price flexibility is cheaper", 1),
    "Discount is modelled as hitting all baseline volume — the standard category-promo assumption.",
], y=1.95, gap=9, size=15)

# ==================================================== 11 · XCHG RESULTS ===
s = slide()
header(s, "Finding 3 · results", "Exchange rate by category")
s.shapes.add_picture(f"{OUT}/fig_exchange_rate.png", Inches(0.4), Inches(2.1), width=Inches(6.8))
rows = [["Category", "Elast.", "Inv $/yr", "Disc.", "Price $/yr", "Rate"]]
for r in findings["exchange_rate"]:
    rows.append([
        r["category"].replace(" & ", " &\n"),
        f"{r['elasticity_base']:.2f}",
        money(r["inv_annual_cost_usd"]),
        f"{r['discount_needed_pct']:.2f}%",
        money(r["price_buffer_cost_usd"]),
        f"{r['exchange_rate']:.2f}×",
    ])
table(s, rows, x=7.5, y=2.0, w=5.55, h=3.2, col_w=[1.5, 0.6, 1.05, 0.65, 1.05, 0.7], fs=10.5)
b = box(s, 7.5, 5.45, 5.6, 1.8)
set_text(b, [
    ("Inventory wins 4 of 5. Frozen Foods is the sole", {"size": 12, "color": GREY}),
    ("break-even (1.23×): high elasticity + low ticket", {"size": 12, "color": GREY}),
    ("price ($5.95) + small SS base vs throughput.", {"size": 12, "color": GREY}),
    ("Beverages worst (0.21×): $36 price × ~2M units.", {"size": 12, "color": GREY}),
], space_after=2)

# ==================================================== 12 · MECHANISM ===
s = slide()
header(s, "Finding 3 · why", "Elastic ≠ cheap-to-buffer-with-price")
left = s.shapes.add_shape(1, Inches(0.7), Inches(2.05), Inches(5.85), Inches(4.6))
right = s.shapes.add_shape(1, Inches(6.75), Inches(2.05), Inches(5.85), Inches(4.6))
for sh in (left, right):
    sh.fill.solid(); sh.fill.fore_color.rgb = CARD
    sh.line.color.rgb = HAIR; sh.line.width = Pt(1)
tf = left.text_frame; tf.word_wrap = True
tf.margin_left = tf.margin_right = Inches(0.3); tf.margin_top = Inches(0.25)
set_text(tf, [
    ("THE BASE-RATE EFFECT", {"size": 13, "color": OCHRE, "bold": True}),
    ("Safety stock is only ~1–1.6% of a category’s annual volume. But the discount needed to "
     "pull in that 1–1.6% is paid on ALL the volume. Carrying cost is paid only on the buffer "
     "units themselves. So even a highly elastic category needs a broad, expensive discount to "
     "match a cheap physical buffer.", {"size": 14, "color": INK}),
], space_after=8, line=1.25)
tf = right.text_frame; tf.word_wrap = True
tf.margin_left = tf.margin_right = Inches(0.3); tf.margin_top = Inches(0.25)
set_text(tf, [
    ("WHERE THE EXCEPTION SHOWS UP", {"size": 13, "color": OCHRE, "bold": True}),
    ("Frozen Foods breaks even: elasticity −3.50, but a low average price ($5.95) and a smaller "
     "safety-stock base relative to throughput mean a small discount goes far. Beverages is the "
     "opposite — $36 average price across ~2M units makes even a 1.1% discount dwarf the cost of "
     "just holding the stock.", {"size": 14, "color": INK}),
], space_after=8, line=1.25)

# ==================================================== 13 · RECOMMEND ===
s = slide()
header(s, "Recommendation", "Buffer-mix starting point per segment")
rows = [["Category", "Rate", "Recommended buffer strategy"]]
reco = {
    "Frozen Foods": "Trial a 25–50% price-buffer blend — closest to parity, highly elastic, low ticket price.",
    "Snacks & Confectionery": "Inventory-led; use price as a secondary lever, mainly on the Value tier.",
    "Grocery": "Inventory-led. Reserve price moves for clearance / write-off avoidance.",
    "Dairy & Bakery": "Inventory-led. Treat Premium and Value tiers as different SKUs — Value does not respond to price.",
    "Beverages": "Inventory-led. Price is the most expensive buffer here by a wide margin.",
}
order = ["Frozen Foods", "Snacks & Confectionery", "Grocery", "Dairy & Bakery", "Beverages"]
xr = {r["category"]: r["exchange_rate"] for r in findings["exchange_rate"]}
for c in order:
    rows.append([c, f"{xr[c]:.2f}×", reco[c]])
table(s, rows, x=0.7, y=2.2, w=11.9, h=3.4, col_w=[2.5, 1.0, 8.4], fs=12.5)
b = box(s, 0.7, 5.9, 11.9, 1.1)
set_text(b, ("All routing is anchored to MEIO’s optimised safety stock. Against the pre-optimisation "
             "policy the inventory side is ~60% larger, which shifts the rates but not the ranking.",
             {"size": 12, "color": GREY}))

# ==================================================== 14 · CAVEATS ===
s = slide()
header(s, "Caveats & next steps", "What would sharpen this")
bullets(s, [
    "Model R² is low (0.008–0.062) — trust the sign / ordering / magnitude, not the point estimates.",
    "Thin price variation: only ~5.7% of lines carry a discount (max 35%). Value-tier Dairy & Bakery "
    "comes back statistically insignificant.",
    "Uniform-discount assumption inflates the price-buffer cost. Targeted / dynamic pricing (the "
    "charter’s agentic layer) would discount only the at-risk slice and move every rate toward price.",
    "One year of actuals — a second year separates seasonal price patterns from true elasticity.",
    "One-shot estimate. The Obsidian research track proposes a self-recalibrating exchange rate via a "
    "multi-armed bandit over buffer-mix arms — a stronger version of the same deliverable.",
    "Exchange rate is anchored to MEIO’s optimised safety stock (not the current policy).",
], y=2.0, gap=12, size=15)
f = box(s, 0.7, 6.7, 11.9, 0.5)
set_text(f, ("Source notebook: Price Elasticity/price_elasticity_analysis.ipynb   ·   "
             "artefacts: Price Elasticity/output/", {"size": 11, "color": GREY}))

prs.save(f"{HERE}/Price_Elasticity_Findings.pptx")
print(f"saved {HERE}/Price_Elasticity_Findings.pptx  ({len(prs.slides)} slides)")
